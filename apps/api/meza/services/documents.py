"""Document ingestion: extraction + rule-based classification (§18)."""

from __future__ import annotations

import re
from pathlib import Path

ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".txt", ".csv", ".png", ".jpg", ".jpeg"}

CLASSIFICATION_RULES: list[tuple[str, list[str]]] = [
    ("CONTRACT", ["договор", "contract", "соглашение"]),
    ("INVOICE", ["счёт", "счет-фактура", "invoice", "накладная"]),
    ("COMMERCIAL_OFFER", ["коммерческое предложение", "commercial offer", "offer"]),
    ("SPECIFICATION", ["спецификац", "specification", "чертёж", "drawing"]),
    ("INSTRUCTION", ["инструкция", "instruction", "manual"]),
    ("REGULATION", ["регламент", "положение", "policy", "regulation"]),
    ("TECH_DOC", ["технич", "technical", "datasheet"]),
    ("REPORT", ["отчёт", "report"]),
]


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}


def ocr_available() -> bool:
    try:
        import pytesseract

        pytesseract.get_tesseract_version()
        return True
    except Exception:  # noqa: BLE001
        return False


def ocr_image(path: Path) -> str:
    """§18: OCR pipeline for scanned/photographed documents, when tesseract is installed on the
    host. Tries Russian+English; falls back to English-only if the Russian language pack isn't
    installed, and to '' (never raises) if tesseract itself is missing."""
    try:
        import pytesseract
        from PIL import Image

        img = Image.open(path)
        try:
            return pytesseract.image_to_string(img, lang="rus+eng")
        except Exception:  # noqa: BLE001
            return pytesseract.image_to_string(img)
    except Exception:  # noqa: BLE001
        return ""


def extract_text(path: Path, ext: str) -> str:
    try:
        if ext == ".txt" or ext == ".csv":
            return path.read_text(errors="ignore")
        if ext == ".pdf":
            from pypdf import PdfReader

            reader = PdfReader(str(path))
            text = "\n".join((page.extract_text() or "") for page in reader.pages)
            return text
        if ext == ".docx":
            from docx import Document as DocxDocument

            d = DocxDocument(str(path))
            return "\n".join(p.text for p in d.paragraphs)
        if ext == ".xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(str(path), read_only=True, data_only=True)
            lines = []
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(" ".join(str(c) for c in row if c is not None))
            return "\n".join(lines)
        if ext in IMAGE_EXTENSIONS:
            return ocr_image(path)
    except Exception:  # noqa: BLE001
        return ""
    return ""


def classify_document(filename: str, text: str) -> tuple[str, float]:
    haystack = f"{filename}\n{text[:5000]}".lower()
    for doc_type, keywords in CLASSIFICATION_RULES:
        for kw in keywords:
            if kw in haystack:
                return doc_type, 0.75
    return "UNCLASSIFIED", 0.0


DATE_RE = re.compile(r"\b(\d{1,2})[./](\d{1,2})[./](\d{4})\b")
AMOUNT_RE = re.compile(r"\b(\d[\d\s]{3,})\s?(тенге|kzt|₸|руб|usd|\$)\b", re.I)


def extract_metadata(text: str) -> dict:
    dates = DATE_RE.findall(text[:20_000])
    amounts = AMOUNT_RE.findall(text[:20_000])
    return {
        "dates_found": [f"{d}.{m}.{y}" for d, m, y in dates[:10]],
        "amounts_found": [f"{a.strip()} {cur}" for a, cur in amounts[:10]],
    }


MAX_CHUNKS_PER_DOCUMENT = 40


def chunk_text(text: str, *, chunk_size: int = 800, overlap: int = 100) -> list[str]:
    """Simple sliding-window chunker (§18 RAG scope: contracts/specs/regulations, not a general
    corpus). Capped at MAX_CHUNKS_PER_DOCUMENT so embedding a huge document doesn't stall the
    upload endpoint on this hardware — a local model embeds one chunk per HTTP round-trip."""
    text = text.strip()
    if not text:
        return []
    chunks = []
    start = 0
    step = max(1, chunk_size - overlap)
    while start < len(text) and len(chunks) < MAX_CHUNKS_PER_DOCUMENT:
        chunks.append(text[start : start + chunk_size])
        start += step
    return chunks
