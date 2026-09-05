"""Import Center (§44): CSV/XLSX/JSON → detect columns → suggest mapping → validate → preview →
import only after explicit confirmation."""

from __future__ import annotations

import csv
import io
import json

from sqlalchemy.ext.asyncio import AsyncSession

from meza.models import ImportJob, Material, Supplier

TARGET_SCHEMAS: dict[str, dict] = {
    "materials": {
        "required": ["sku", "name", "unit"],
        "optional": ["category", "min_stock", "reorder_quantity", "lead_time_days", "unit_cost"],
    },
    "suppliers": {"required": ["name"], "optional": ["category", "contact", "avg_lead_time_days"]},
}


def detect_columns(file_type: str, content: bytes) -> list[str]:
    if file_type == "csv":
        text = content.decode("utf-8-sig", errors="ignore")
        reader = csv.reader(io.StringIO(text))
        return next(reader, [])
    if file_type == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True)
        ws = wb.worksheets[0]
        row = next(ws.iter_rows(values_only=True), ())
        return [str(c) for c in row if c is not None]
    if file_type == "json":
        data = json.loads(content)
        if isinstance(data, list) and data:
            return list(data[0].keys())
    return []


def suggest_mapping(columns: list[str], target_entity: str) -> dict:
    schema = TARGET_SCHEMAS.get(target_entity, {})
    fields = schema.get("required", []) + schema.get("optional", [])
    mapping = {}
    for field in fields:
        for col in columns:
            if col.strip().lower().replace(" ", "_") == field:
                mapping[field] = col
                break
    return mapping


def read_rows(file_type: str, content: bytes) -> list[dict]:
    if file_type == "csv":
        text = content.decode("utf-8-sig", errors="ignore")
        return list(csv.DictReader(io.StringIO(text)))
    if file_type == "xlsx":
        from openpyxl import load_workbook

        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows_iter = ws.iter_rows(values_only=True)
        header = list(next(rows_iter))
        out = []
        for row in rows_iter:
            out.append({header[i]: row[i] for i in range(len(header)) if i < len(row)})
        return out
    if file_type == "json":
        data = json.loads(content)
        return data if isinstance(data, list) else [data]
    return []


def validate_rows(rows: list[dict], mapping: dict, target_entity: str) -> dict:
    schema = TARGET_SCHEMAS.get(target_entity, {})
    required = schema.get("required", [])
    errors = []
    for i, row in enumerate(rows):
        for field in required:
            col = mapping.get(field)
            if not col or row.get(col) in (None, ""):
                errors.append({"row": i, "field": field, "error": "missing required value"})
    return {"valid": len(errors) == 0, "errors": errors[:50], "row_count": len(rows)}


async def apply_import(db: AsyncSession, job: ImportJob) -> int:
    rows = read_rows(job.file_type, open(job.stored_path, "rb").read())
    mapping = job.mapping
    imported = 0
    if job.target_entity == "materials":
        for row in rows:
            sku = row.get(mapping.get("sku", ""))
            if not sku:
                continue
            m = Material(
                sku=str(sku), name=str(row.get(mapping.get("name", ""), "")),
                unit=str(row.get(mapping.get("unit", ""), "kg")),
                category=str(row.get(mapping.get("category", ""), "") or ""),
                min_stock=float(row.get(mapping.get("min_stock", ""), 0) or 0),
                reorder_quantity=float(row.get(mapping.get("reorder_quantity", ""), 0) or 0),
                unit_cost=float(row.get(mapping.get("unit_cost", ""), 0) or 0),
            )
            db.add(m)
            imported += 1
    elif job.target_entity == "suppliers":
        for row in rows:
            name = row.get(mapping.get("name", ""))
            if not name:
                continue
            db.add(Supplier(name=str(name), category=str(row.get(mapping.get("category", ""), "") or "")))
            imported += 1
    await db.flush()
    return imported
